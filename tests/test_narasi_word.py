"""Narration, table and Word helpers. Gemini is never called."""

import pandas as pd
import pytest

from Main.config import cfg
from Main import narasi
from Main.utils import count_points
from Main.word import _merge_runs, report_filename


@pytest.fixture
def set_cfg():
    saved = (cfg.peta, cfg.tipe, cfg.skala)

    def _set(peta, tipe, skala):
        cfg.peta, cfg.tipe, cfg.skala = peta, tipe, skala
    yield _set
    cfg.peta, cfg.tipe, cfg.skala = saved


def _counts(**kw):
    d = dict(kw)
    d['total'] = sum(kw.values())
    return d


# ---- count_points ----------------------------------------------------------

def test_count_points_ch_edges_and_total(set_cfg):
    set_cfg('Prakiraan', 'Curah Hujan', 'Bulanan')
    df = pd.DataFrame({'VAL': [0, 99, 100, 300, 499, 500, 900, None]})
    c = count_points(df, 'VAL', levels=None)
    assert c == {'Rendah': 2, 'Menengah': 1, 'Tinggi': 2, 'Sangat Tinggi': 2, 'total': 7}


def test_count_points_verifikasi_not_binned_as_ch(set_cfg):
    # Verifikasi runs with tipe 'Curah Hujan'; it used to fall into the CH bins.
    set_cfg('Verifikasi', 'Curah Hujan', 'Bulanan')
    c = count_points(pd.DataFrame({'v': [0, 0, 1]}), 'v', levels=[0, 1])
    assert c == {'Tidak Sesuai': 2, 'Sesuai': 1, 'total': 3}


def test_count_points_bias_counts_negatives(set_cfg):
    set_cfg('Bias', 'Curah Hujan', 'Bulanan')
    c = count_points(pd.DataFrame({'b': [-300, -50, 0, 30, 150]}), 'b', levels=None)
    assert c['total'] == 5
    assert c['Negatif Tinggi'] == 1 and c['Positif Tinggi'] == 1


def test_count_points_spi_boundaries_match_labels(set_cfg):
    # Dry classes include the upper bound (<=-2,0), wet ones the lower (>=2,0)
    set_cfg('Analisis', 'SPI', 'Bulanan')
    c = count_points(pd.DataFrame({'SPI': [-2.0, -1.5, -1.0, 0, 1.0, 1.5, 2.0]}), 'SPI', levels=None)
    assert c == {'Sangat Kering': 1, 'Kering': 1, 'Agak Kering': 1, 'Normal': 1,
                 'Agak Basah': 1, 'Basah': 1, 'Sangat Basah': 1, 'total': 7}


# ---- grouping, narration, table -------------------------------------------

def _map_data():
    return {
        'peta': 'Prakiraan', 'tipe': 'Curah Hujan', 'skala': 'Dasarian',
        'year': 2026, 'month': 2, 'dasarian': 3,
        'nama_wilayah': 'Provinsi PAPUA BARAT DAN PAPUA BARAT DAYA',
        'province_data': {
            'PAPUA BARAT': _counts(Rendah=1, Menengah=6, Tinggi=3, **{'Sangat Tinggi': 0}),
            'PAPUA BARAT DAYA': _counts(Rendah=2, Menengah=5, Tinggi=3, **{'Sangat Tinggi': 0}),
        },
        'kabupaten_data': {
            'MANOKWARI': _counts(Rendah=0, Menengah=9, Tinggi=11, **{'Sangat Tinggi': 0}),  # 45/55 near-tie
            'SORONG': _counts(Rendah=2, Menengah=18, Tinggi=0, **{'Sangat Tinggi': 0}),  # Rendah 10% extreme
            'MAYBRAT': _counts(Rendah=0, Menengah=10, Tinggi=0, **{'Sangat Tinggi': 0}),
            'EMPTY': _counts(Rendah=0, Menengah=0, Tinggi=0, **{'Sangat Tinggi': 0}),
        },
    }


def test_group_kabupaten_rules():
    md = _map_data()
    g = narasi.group_kabupaten(md['kabupaten_data'], narasi._categories(md))
    assert g['Tinggi'] == ['Manokwari']
    assert g['Menengah'] == ['Manokwari', 'Maybrat', 'Sorong']  # Manokwari: near-tie
    assert g['Rendah'] == ['Sorong']                            # extreme >= 5%
    assert g['Sangat Tinggi'] == []


def test_template_narration_exact_and_consistent_with_table():
    md = _map_data()
    text = narasi.get_analysis(md, use_llm=False)
    assert text.startswith(
        "Prakiraan Curah Hujan Dasarian III Februari 2026 (21 - 28 Februari 2026) "
        "di Provinsi Papua Barat dan Papua Barat Daya umumnya didominasi curah hujan "
        "Menengah (50-150 mm/das) sebesar 60.0% di Papua Barat dan 50.0% di Papua Barat Daya."
    )
    assert "Rendah (0-50 mm/das) diprakirakan terjadi di Kab. Sorong" in text
    assert "Tinggi (150-300 mm/das) diprakirakan terjadi di Kab. Manokwari" in text

    table = narasi.build_table_data(md)
    rows = dict(table['rows'])
    assert rows['Rendah (0-50 mm/das)'] == 'Sorong'
    assert rows['Tinggi (150-300 mm/das)'] == 'Manokwari'
    assert rows['Sangat Tinggi (>300 mm/das)'] == '-'


def test_different_province_dominants_still_name_kabupaten():
    md = _map_data()
    md['province_data']['PAPUA BARAT'] = _counts(Rendah=0, Menengah=3, Tinggi=7, **{'Sangat Tinggi': 0})
    text = narasi.get_analysis(md, use_llm=False)
    assert "sementara Provinsi Papua Barat Daya didominasi" in text
    assert "Tinggi (150-300 mm/das) diprakirakan terjadi di Kab. Manokwari" in text
    assert "Menengah (50-150 mm/das) diprakirakan terjadi di Kab. Manokwari, Maybrat, dan Sorong" in text


class _FakeClient:
    def __init__(self, text):
        resp = type('R', (), {'text': text})()
        self.models = type('M', (), {'generate_content': lambda _s, **kw: resp})()


@pytest.mark.parametrize('llm_text, expect_llm', [
    ("**Curah hujan Tinggi (150-300 mm/das) diprakirakan terjadi di Kab. Manokwari, "
     "sedangkan Rendah (0-50 mm/das) di Kab. Sorong.**", True),
    ("Curah hujan Tinggi diprakirakan terjadi di Kab. Manokwari.", False),   # drops Sorong
    ("Tinggi 40% di Manokwari dan Sorong.", False),                          # percentage
    ("Manokwari dan Sorong.", False),                                        # no categories
    (None, False),
])
def test_llm_output_checked(monkeypatch, llm_text, expect_llm):
    monkeypatch.setattr(narasi, '_gemini_client', lambda: (None, None, _FakeClient(llm_text)))
    text = narasi.get_analysis(_map_data())
    if expect_llm:
        assert "sedangkan Rendah (0-50 mm/das) di Kab. Sorong." in text
        assert '*' not in text
    else:
        assert "Rendah (0-50 mm/das) diprakirakan terjadi di Kab. Sorong" in text


def test_percent_inside_category_range_is_allowed():
    items = [('Atas Normal', '>115%', ['Manokwari'])]
    assert narasi._valid_kabupaten_text("Sifat hujan Atas Normal (>115%) tercatat di Kab. Manokwari.", items)
    assert not narasi._valid_kabupaten_text("Atas Normal 40% di Kab. Manokwari.", items)


def test_verifikasi_metrics_and_kabupaten_level():
    md = {
        'peta': 'Verifikasi', 'tipe': 'Curah Hujan', 'skala': 'Bulanan',
        'year': 2026, 'month': 9, 'dasarian': 1,
        'nama_wilayah': 'Kabupaten SORONG',
        'province_data': {'PAPUA BARAT DAYA': _counts(**{'Tidak Sesuai': 3, 'Sesuai': 1})},
        'kabupaten_data': {'SORONG': _counts(**{'Tidak Sesuai': 3, 'Sesuai': 1})},
        'accuracy': 0.25, 'hss': 0.012, 'pss': 0.12,
    }
    text = narasi.get_analysis(md, use_llm=False)
    assert "di Kabupaten Sorong didominasi kategori Tidak Sesuai sebesar 75.0%." in text
    assert text.endswith("Akurasi keseluruhan (PC) mencapai 25.00%, HSS sebesar 1.20%, dan PSS sebesar 12.00%.")
    assert narasi.build_table_data(md) is None


# ---- word helpers ----------------------------------------------------------

def test_merge_runs_respects_parent_column():
    rows = [['A', 'x', '1'], ['A', 'x', '2'], ['A', 'y', '3'], ['B', 'y', '4'], ['B', 'y', '5']]
    assert sorted(_merge_runs(rows, 2)) == [(0, 0, 2), (0, 3, 4), (1, 0, 1), (1, 3, 4)]


def test_report_filename_unique_and_safe():
    a = report_filename('Prakiraan Curah Hujan', 'Bulan Oktober Dasarian III 2026')
    b = report_filename('Prakiraan Curah Hujan', 'Bulan Oktober 2026')
    assert a != b and a.endswith('.docx')
    assert report_filename('A/B', 'C:D') == 'Laporan AB CD.docx'


def test_hth_table_cells_merge_in_docx(tmp_path):
    from docx import Document
    from docxtpl import DocxTemplate
    from Main.word import _build_table_subdoc

    tpl = tmp_path / 't.docx'
    d = Document()
    d.add_paragraph('{{p text2 }}')
    d.save(tpl)

    doc = DocxTemplate(str(tpl))
    data = {'columns': ['Provinsi', 'Kabupaten', 'Indeks HTH'],
            'rows': [['P', 'K1', 'a'], ['P', 'K1', 'b'], ['P', 'K2', 'c']],
            'merge_cols': 2}
    doc.render({'text2': _build_table_subdoc(doc, data)})
    out = tmp_path / 'o.docx'
    doc.save(out)

    table = Document(out).tables[0]
    col0 = [c.text for c in table.columns[0].cells]
    col1 = [c.text for c in table.columns[1].cells]
    assert col0 == ['Provinsi', 'P', 'P', 'P']       # one merged cell, read back per row
    assert table.cell(1, 0)._tc is table.cell(3, 0)._tc
    assert col1 == ['Kabupaten', 'K1', 'K1', 'K2']
    assert table.cell(1, 1)._tc is table.cell(2, 1)._tc
    assert table.cell(2, 1)._tc is not table.cell(3, 1)._tc


def test_arrange_word_keeps_angle_brackets(tmp_path, monkeypatch):
    from docx import Document
    from PIL import Image
    import Main.word as word

    d = Document()
    d.add_paragraph('{{ title1 }}')
    d.add_paragraph('{{ image1 }}')
    d.add_paragraph('{{ text1 }}')
    d.add_paragraph('{{p text2 }}')
    d.save(tmp_path / 'template_doc.docx')
    monkeypatch.setattr(word, 'CACHE_DIR', str(tmp_path))
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(word, 'get_full_narration', lambda md: 'Bawah Normal (<85%) & Atas Normal (>115%)')
    monkeypatch.setattr(word, 'build_table_data',
                        lambda md: {'columns': ['Sifat Hujan', 'Daerah'], 'rows': [['Bawah Normal (<85%)', 'Sorong']]})

    md = {'peta': 'Prakiraan', 'tipe': 'Sifat Hujan', 'skala': 'Bulanan', 'year': 2026, 'month': 11,
          'image': Image.new('RGB', (20, 20))}
    out = word.arrange_word(md)
    doc = Document(out)
    texts = [p.text for p in doc.paragraphs]
    assert 'Bawah Normal (<85%) & Atas Normal (>115%)' in texts
    assert doc.tables[0].cell(1, 0).text == 'Bawah Normal (<85%)'
    assert doc.inline_shapes  # image survived autoescape


def test_long_merged_table_builds(tmp_path):
    # cell.merge() used to hit RecursionError around 1000 merged rows.
    from docxtpl import DocxTemplate
    from docx import Document
    from Main.word import _build_table_subdoc

    tpl = tmp_path / 't.docx'
    d = Document()
    d.add_paragraph('x')
    d.save(tpl)
    doc = DocxTemplate(str(tpl))
    rows = [['P', 'K', str(i)] for i in range(1500)]
    sd = _build_table_subdoc(doc, {'columns': ['a', 'b', 'c'], 'rows': rows, 'merge_cols': 2})
    assert len(sd.tables[0].rows) == 1501
